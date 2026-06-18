import customtkinter as ctk
from tkinter import ttk, messagebox, filedialog
import pandas as pd
from database.db_manager import (
    get_all_equipment,
    search_equipment,
    delete_equipment,
    get_statistics,
    filter_by_status
)
from theme import COLORS
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import PatternFill, Font


class MainWindow:
    def __init__(self, root):
        self.root = root
        self.root.title("IT Equipment Manager")
        self.root.geometry("1400x800")
        self.root.configure(fg_color=COLORS["bg"])

        self.setup_layout()
        self.setup_sidebar()
        self.setup_main_area()
        self.load_data()

    def setup_layout(self):
        self.sidebar = ctk.CTkFrame(
            self.root,
            width=220,
            corner_radius=0,
            fg_color=COLORS["sidebar_bg"]
        )
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)  

        self.main = ctk.CTkFrame(
            self.root,
            corner_radius=0,
            fg_color=COLORS["bg"]
        )
        self.main.pack(side="left", fill="both", expand=True)


    def setup_sidebar(self):
        logo_frame = ctk.CTkFrame(
            self.sidebar,
            fg_color=COLORS["primary"],
            corner_radius=0,
            height=70
        )
        logo_frame.pack(fill="x")
        logo_frame.pack_propagate(False)

        ctk.CTkLabel(
            logo_frame,
            text="⚙  IT Manager",
            font=ctk.CTkFont("Segoe UI", 16, "bold"),
            text_color="white"
        ).place(relx=0.5, rely=0.5, anchor="center")

        ctk.CTkFrame(
            self.sidebar,
            height=1,
            fg_color=COLORS["border"]
        ).pack(fill="x", pady=(0, 10))

        nav_items = [
            ("Dashboard",  "📊",  "dashboard"),
            ("Equipment",  "💻",  "equipment"),
            ("Add New",    "➕",  "add"),
            ("Export",     "📤",  "export"),
        ]

        self.nav_buttons = {}

        for label, icon, section in nav_items:
            btn = ctk.CTkButton(
                self.sidebar,
                text=f"  {icon}  {label}",
                anchor="w",
                height=45,
                corner_radius=8,
                fg_color="transparent",
                hover_color=COLORS["sidebar_hover"],
                text_color=COLORS["text_sidebar"],
                font=ctk.CTkFont("Segoe UI", 13),
                command=lambda s=section: self.navigate(s)
            )
            btn.pack(fill="x", padx=10, pady=3)
            self.nav_buttons[section] = btn

        self.set_active_nav("equipment")

        ctk.CTkLabel(
            self.sidebar,
            text="v1.0  • © 2026",
            font=ctk.CTkFont("Segoe UI", 10),
            text_color=COLORS["text_dim"]
        ).pack(side="bottom", pady=15)

    def set_active_nav(self, section):
        for key, btn in self.nav_buttons.items():
            if key == section:
                btn.configure(
                    fg_color=COLORS["primary"],
                    text_color="white"
                )
            else:
                btn.configure(
                    fg_color="transparent",
                    text_color=COLORS["text_sidebar"]
                )

    def navigate(self, section):
        self.set_active_nav(section)

        if section == "add":
            self.open_add_form()

            self.set_active_nav("equipment")
        elif section == "export":
            self.export_to_excel()
            self.set_active_nav("equipment")
        elif section == "dashboard":
            self.show_dashboard_view()
        elif section == "equipment":
            self.show_equipment_view()


    def setup_main_area(self):
        topbar = ctk.CTkFrame(
            self.main,
            height=60,
            corner_radius=0,
            fg_color=COLORS["surface"]
        )
        topbar.pack(fill="x")
        topbar.pack_propagate(False)

        self.search_var = ctk.StringVar()
        ctk.CTkEntry(
            topbar,
            textvariable=self.search_var,
            placeholder_text="🔍  Search equipment...",
            width=300,
            height=36,
            corner_radius=8,
            fg_color=COLORS["entry_bg"],
            border_color=COLORS["border"],
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 12)
        ).place(x=20, rely=0.5, anchor="w")

        ctk.CTkButton(
            topbar,
            text="Search",
            width=80,
            height=36,
            corner_radius=8,
            fg_color=COLORS["primary"],
            hover_color=COLORS["primary_hover"],
            font=ctk.CTkFont("Segoe UI", 12),
            command=self.perform_search
        ).place(x=330, rely=0.5, anchor="w")

        ctk.CTkButton(
            topbar,
            text="Clear",
            width=70,
            height=36,
            corner_radius=8,
            fg_color=COLORS["surface_light"],
            hover_color=COLORS["sidebar_hover"],
            font=ctk.CTkFont("Segoe UI", 12),
            command=self.clear_search
        ).place(x=420, rely=0.5, anchor="w")

        ctk.CTkLabel(
            topbar,
            text="Status:",
            text_color=COLORS["text_dim"],
            font=ctk.CTkFont("Segoe UI", 12)
        ).place(x=510, rely=0.5, anchor="w")

        self.status_filter = ctk.StringVar(value="All")
        ctk.CTkOptionMenu(
            topbar,
            variable=self.status_filter,
            values=["All", "Active", "Maintenance", "In Storage"],
            width=130,
            height=36,
            corner_radius=8,
            fg_color=COLORS["entry_bg"],
            button_color=COLORS["primary"],
            button_hover_color=COLORS["primary_hover"],
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 12),
            command=self.apply_status_filter
        ).place(x=560, rely=0.5, anchor="w")

        ctk.CTkButton(
            topbar,
            text="✎  Edit",
            width=90,
            height=36,
            corner_radius=8,
            fg_color=COLORS["primary"],
            hover_color=COLORS["primary_hover"],
            font=ctk.CTkFont("Segoe UI", 12, "bold"),
            command=self.open_edit_form
        ).place(relx=0.82, rely=0.5, anchor="w")

        ctk.CTkButton(
            topbar,
            text="✕  Delete",
            width=100,
            height=36,
            corner_radius=8,
            fg_color=COLORS["danger"],
            hover_color="#c62828",
            font=ctk.CTkFont("Segoe UI", 12, "bold"),
            command=self.delete_selected
        ).place(relx=0.91, rely=0.5, anchor="w")

        self.content = ctk.CTkFrame(
            self.main,
            corner_radius=0,
            fg_color=COLORS["bg"]
        )
        self.content.pack(fill="both", expand=True)

        self.setup_dashboard_cards()
        self.setup_table()


        self.show_equipment_view()

    def setup_dashboard_cards(self):
        self.dashboard_frame = ctk.CTkFrame(
            self.content,
            fg_color=COLORS["bg"],
            corner_radius=0
        )

        cards_row = ctk.CTkFrame(
            self.dashboard_frame,
            fg_color="transparent"
        )
        cards_row.pack(fill="x", padx=20, pady=20)

        card_data = [
            ("Total Assets",  "total",       COLORS["primary"],     "💻"),
            ("Active",        "active",       COLORS["success"],     "✅"),
            ("Maintenance",   "maintenance",  COLORS["maintenance"], "🔧"),
            ("In Storage",    "storage",      COLORS["surface_light"],"📦"),
        ]

        self.stat_labels = {}

        for title, key, color, icon in card_data:
            card = ctk.CTkFrame(
                cards_row,
                fg_color=color,
                corner_radius=12,
                height=100
            )
            card.pack(side="left", expand=True, fill="x", padx=8)
            card.pack_propagate(False)

            ctk.CTkLabel(
                card,
                text=f"{icon}  {title}",
                font=ctk.CTkFont("Segoe UI", 12),
                text_color="white"
            ).place(relx=0.5, rely=0.3, anchor="center")

            val = ctk.CTkLabel(
                card,
                text="0",
                font=ctk.CTkFont("Segoe UI", 28, "bold"),
                text_color="white"
            )
            val.place(relx=0.5, rely=0.68, anchor="center")
            self.stat_labels[key] = val

    def show_dashboard_view(self):
        self.table_frame.pack_forget()
        self.dashboard_frame.pack(fill="both", expand=True)
        self.refresh_dashboard()

    def show_equipment_view(self):
        self.dashboard_frame.pack_forget()
        self.table_frame.pack(fill="both", expand=True)

    def refresh_dashboard(self):
        stats = get_statistics()
        for key, label in self.stat_labels.items():
            label.configure(text=str(stats[key]))


    def setup_table(self):
        self.table_frame = ctk.CTkFrame(
            self.content,
            fg_color=COLORS["bg"],
            corner_radius=0
        )

        columns = (
            "ID", "Tag", "Type", "Brand", "Model",
            "Serial Number", "Status", "Department",
            "Purchase Date", "Specs"
        )

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview",
            background=COLORS["surface"],
            foreground=COLORS["text"],
            fieldbackground=COLORS["surface"],
            rowheight=32,
            font=("Segoe UI", 10),
            borderwidth=0
        )
        style.configure("Treeview.Heading",
            background=COLORS["sidebar_bg"],
            foreground=COLORS["text_dim"],
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            borderwidth=0
        )
        style.map("Treeview",
            background=[("selected", COLORS["primary"])],
            foreground=[("selected", "white")]
        )
        style.layout("Treeview", [
            ('Treeview.treearea', {'sticky': 'nswe'})
        ])

        self.table = ttk.Treeview(
            self.table_frame,
            columns=columns,
            show="headings",
            style="Treeview"
        )

        col_widths = {
            "ID": 45, "Tag": 90, "Type": 95, "Brand": 90,
            "Model": 115, "Serial Number": 120, "Status": 90,
            "Department": 110, "Purchase Date": 110, "Specs": 160
        }

        for col in columns:
            self.table.heading(
                col,
                text=col,
                command=lambda c=col: self.sort_column(c, False)
            )
            self.table.column(
                col,
                width=col_widths.get(col, 100),
                anchor="center" if col == "ID" else "w"
            )

        self.table.tag_configure(
            "odd",  background=COLORS["surface"]
        )
        self.table.tag_configure(
            "even", background=COLORS["surface_light"]
        )

        scrollbar = ctk.CTkScrollbar(
            self.table_frame,
            command=self.table.yview,
            fg_color=COLORS["surface"],
            button_color=COLORS["primary"],
            button_hover_color=COLORS["primary_hover"]
        )
        self.table.configure(yscrollcommand=scrollbar.set)

        self.table.pack(
            side="left", fill="both",
            expand=True, padx=(15, 0), pady=10
        )
        self.table.bind("<Double-1>", self.on_double_click)
        scrollbar.pack(side="right", fill="y", pady=10, padx=(0, 5))
        
    def on_double_click(self, event):
        region = self.table.identify_region(event.x, event.y)
        if region == "cell":
            self.open_edit_form()


    def load_data(self, data=None):
        for row in self.table.get_children():
            self.table.delete(row)

        rows = data if data is not None else get_all_equipment()

        for i, row in enumerate(rows):
            tag = "even" if i % 2 == 0 else "odd"
            self.table.insert("", "end", tags=(tag,), values=(
                row[0], row[1], row[2], row[3], row[4],
                row[5], row[6], row[7], row[8], row[9],
            ))
        self.refresh_dashboard()

    def perform_search(self):
        keyword = self.search_var.get().strip()
        if keyword:
            self.load_data(search_equipment(keyword))

    def clear_search(self):
        self.search_var.set("")
        self.status_filter.set("All")
        self.load_data()

    def apply_status_filter(self, selected):
        if selected == "All":
            self.load_data()
        else:
            self.load_data(filter_by_status(selected))

    def get_selected_id(self):
        selected = self.table.selection()
        if not selected:
            messagebox.showwarning(
                "No Selection",
                "Please select an equipment row first."
            )
            return None
        return self.table.item(selected[0])["values"][0]

    def delete_selected(self):
        equipment_id = self.get_selected_id()
        if equipment_id is None:
            return
        if messagebox.askyesno(
            "Confirm Delete",
            "Are you sure you want to delete this equipment?"
        ):
            delete_equipment(equipment_id)
            self.load_data()

    def open_add_form(self):
        from gui.add_asset_form import AddEquipmentForm
        AddEquipmentForm(self.root, self.load_data)

    def open_edit_form(self):
        from gui.edit_asset_form import EditEquipmentForm
        equipment_id = self.get_selected_id()
        if equipment_id is None:
            return
        EditEquipmentForm(self.root, equipment_id, self.load_data)


    def export_to_excel(self):
        rows = get_all_equipment()
        if not rows:
            messagebox.showwarning(
                "No Data",
                "There is no equipment to export."
            )
            return

        columns = [
            "ID", "Asset Tag", "Type", "Brand", "Model",
            "Serial Number", "Status", "Department",
            "Purchase Date", "Specs"
        ]

        df = pd.DataFrame(rows, columns=columns)

        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="equipment_report.xlsx",
            title="Save Excel Report"
        )

        if not file_path:
            return

        try:
            with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Equipment")
                ws = writer.sheets["Equipment"]

                ws.insert_rows(1)
                ws["A1"] = "IT Equipment Report"
                ws["A1"].font = Font(size=16, bold=True)

                table = Table(
                    displayName="EquipmentTable",
                    ref=f"A2:J{ws.max_row}"
                )
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showRowStripes=True
                )
                ws.add_table(table)

                for col in ws.columns:
                    max_len = max(
                        (len(str(c.value)) for c in col if c.value),
                        default=10
                    )
                    ws.column_dimensions[
                        col[0].column_letter
                    ].width = max_len + 4

                ws.freeze_panes = "A3"

                fills = {
                    "Active":      PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
                    "Maintenance": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
                    "In Storage":  PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
                }

                for row in range(3, ws.max_row + 1):
                    cell = ws[f"G{row}"]
                    if cell.value in fills:
                        cell.fill = fills[cell.value]

            messagebox.showinfo(
                "Success",
                f"Exported successfully to:\n{file_path}"
            )
        except Exception as e:
            messagebox.showerror("Error", f"Export failed.\n{e}")
            
    def sort_column(self, col, reverse):
        rows = [
            (self.table.set(row_id, col), row_id)
            for row_id in self.table.get_children("")
        ]

        try:
            rows.sort(key=lambda x: float(x[0]), reverse=reverse)
        except ValueError:
            rows.sort(key=lambda x: x[0].lower(), reverse=reverse)

        for index, (_, row_id) in enumerate(rows):
            self.table.move(row_id, "", index)

            tag = "even" if index % 2 == 0 else "odd"
            self.table.item(row_id, tags=(tag,))

        self.table.heading(
            col,
            command=lambda: self.sort_column(col, not reverse)
        )